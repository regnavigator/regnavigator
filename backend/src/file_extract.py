"""업로드 파일 텍스트 추출 — 파일 업로드 v1 (2026-08-05).

설계 원칙:
- **원본을 저장하지 않는다.** 추출 텍스트만 반환하고 바이트는 버린다 —
  개인정보 보관·파기 부담을 v1 에서 만들지 않는다. 추출 텍스트는 이후
  /chat 첨부로 들어가 기존 PII 게이트·at-rest 마스킹을 그대로 통과한다.
- 지원: PDF(텍스트 레이어) · DOCX · HWPX · TXT/MD.
- 구형 HWP(OLE)·DOC 은 **한국어 안내와 함께 거절** — 몰래 미리보기 텍스트만
  뽑으면 잘린 문서를 전체라고 믿게 된다(silent truncation 금지).
- 스캔 PDF(텍스트 레이어 없음)도 같은 이유로 거절.
"""

from __future__ import annotations

import io
import logging
import os
import re
import zipfile
from dataclasses import dataclass

_log = logging.getLogger(__name__)

# OLE 컨테이너 매직 — 구형 .hwp/.doc 판별
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_ZIP_MAGIC = b"PK\x03\x04"


class UnsupportedFileType(Exception):
    """지원하지 않는 형식 — message 는 사용자에게 그대로 보여줄 한국어."""


class EmptyExtract(Exception):
    """파일은 읽었으나 텍스트가 없음(스캔 PDF 등) — 한국어 안내."""


@dataclass
class ExtractResult:
    text: str
    kind: str          # pdf | docx | hwpx | txt
    truncated: bool


def max_extract_chars() -> int:
    """추출 텍스트 상한(문자). 초과분은 잘라내고 truncated=True 로 알린다.

    120,000자(2026-09-14, 종전 30,000) — 임상시험계획서 50쪽 안팎을 통째로 받기
    위해. 종전 상한의 근거였던 PII 게이트 처리량은 NER 교체(8/14) 후 실측
    120,000자 1.2s 로 병목이 아니다. `routes.chat` 의 파일당 방어 상한이 이 값을
    따른다.
    """
    try:
        return max(1000, int((os.getenv("FILE_MAX_EXTRACT_CHARS") or "").strip() or "120000"))
    except ValueError:
        return 120000


def extract_text(filename: str, data: bytes) -> ExtractResult:
    """파일 바이트 → 텍스트. 실패는 UnsupportedFileType / EmptyExtract."""
    ext = (filename.rsplit(".", 1)[-1] if "." in filename else "").lower()

    if ext == "hwp" and data[:4] == _ZIP_MAGIC:
        ext = "hwpx"  # 확장자만 .hwp 로 저장된 hwpx

    if ext == "pdf":
        text, kind = _extract_pdf(data), "pdf"
    elif ext == "docx":
        text, kind = _extract_docx(data), "docx"
    elif ext == "hwpx":
        text, kind = _extract_hwpx(data), "hwpx"
    elif ext in ("txt", "md", "csv"):
        text, kind = _decode_text(data), "txt"
    elif ext in ("hwp", "doc"):
        raise UnsupportedFileType(
            "구형 HWP/DOC 형식은 지원하지 않습니다. 한글/워드에서 "
            "HWPX·DOCX 또는 PDF로 저장한 뒤 올려주세요."
        )
    else:
        raise UnsupportedFileType(
            f"지원하지 않는 형식입니다({ext or '확장자 없음'}). "
            "PDF·DOCX·HWPX·TXT 파일을 올려주세요."
        )

    text = _normalize(text)
    if not text.strip():
        raise EmptyExtract(
            "텍스트를 추출할 수 없습니다. 스캔(이미지) PDF는 지원하지 않습니다 — "
            "텍스트가 포함된 원본 파일로 올려주세요."
        )

    cap = max_extract_chars()
    if len(text) > cap:
        return ExtractResult(text=text[:cap], kind=kind, truncated=True)
    return ExtractResult(text=text, kind=kind, truncated=False)


# ---------------------------------------------------------------------------
# 형식별 추출
# ---------------------------------------------------------------------------

def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise UnsupportedFileType("암호가 걸린 PDF는 지원하지 않습니다. 암호를 해제한 뒤 올려주세요.")
        pages = [(p.extract_text() or "") for p in reader.pages]
    except (UnsupportedFileType, EmptyExtract):
        raise
    except Exception as e:  # 손상 파일 등
        _log.warning("pdf extract failed: %s", e)
        raise UnsupportedFileType("PDF를 읽을 수 없습니다. 파일이 손상되지 않았는지 확인해 주세요.")
    return "\n\n".join(pages)


def _extract_docx(data: bytes) -> str:
    from docx import Document
    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:
        _log.warning("docx extract failed: %s", e)
        raise UnsupportedFileType("DOCX를 읽을 수 없습니다. 파일이 손상되지 않았는지 확인해 주세요.")
    parts: list[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _extract_hwpx(data: bytes) -> str:
    """HWPX = zip + OWPML XML. 본문 텍스트는 <hp:t> 요소."""
    import xml.etree.ElementTree as ET
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
        sections = sorted(
            n for n in zf.namelist()
            if n.startswith("Contents/section") and n.endswith(".xml")
        )
        if not sections:
            raise UnsupportedFileType("HWPX 본문(section)을 찾을 수 없습니다.")
        out: list[str] = []
        for name in sections:
            root = ET.fromstring(zf.read(name))
            for el in root.iter():
                tag = el.tag.rsplit("}", 1)[-1]
                if tag == "t" and el.text:
                    out.append(el.text)
                elif tag == "p":
                    out.append("\n")
        return "".join(out)
    except UnsupportedFileType:
        raise
    except Exception as e:
        _log.warning("hwpx extract failed: %s", e)
        raise UnsupportedFileType("HWPX를 읽을 수 없습니다. 파일이 손상되지 않았는지 확인해 주세요.")


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8", "cp949"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


_WS_RE = re.compile(r"[ \t ]+")
_NL_RE = re.compile(r"\n{3,}")


def _normalize(text: str) -> str:
    """공백 정리 — 토큰 낭비를 줄이되 문단 구조(빈 줄 1개)는 유지."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _WS_RE.sub(" ", text)
    text = _NL_RE.sub("\n\n", text)
    return text.strip()
